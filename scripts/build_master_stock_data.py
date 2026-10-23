#!/usr/bin/env python3
"""Clean Gaurav's raw-material / premix / blend tabs into structured data for the office app.

Reads an .xlsx export of the live Google Sheet "Technofluid-Product-Master-GRADES"
(id 1P2F2T0kqQtA_8qwuzpPq7rtRic02jMTP9yaEpBEhPLw) and writes office/data/master/:

  raw-materials.json   Sheet1 rows 20+  → stock items, kind raw_material
  consumables.json     Sheet1 rows 2–19 → stock items, kind consumable (cans, caps, cartons)
  premixes.json        Sheet2           → "VI" premixes (base oil + Lubimax 141 polymer)
  blends.json          Sheet3           → house base blends (r5, sn, g90, va17.3, h68, tq …)
  product-recipes.json main tab cols F–L, one recipe per stock pool (product + grade)
  resolution.json      every ingredient code used anywhere → what it resolves to (or not)
  source/*.csv         the three tabs exactly as exported, for traceability

Nothing is invented: names keep Gaurav's wording as an alias, every assumption is listed
in `flags`, and codes we cannot resolve are reported instead of guessed.

Usage:  python3 scripts/build_master_stock_data.py <master.xlsx> [out_dir]
Requires: openpyxl
"""
from __future__ import annotations

import csv
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent

# Plain spelling fixes applied to display names only. Gaurav's original text is always kept.
TYPO_FIXES = {
    "caton": "carton",
    "whie": "white",
    "concenterate": "concentrate",
    "fluis": "fluid",
}

# Aliases we infer but Gaurav has not confirmed — applied, and listed as flags.
ASSUMED_ALIASES = {
    "nvi10kg": ("n4vi10kg", "Sheet3 'nh' uses nvi10kg; Sheet2 defines N4vi10kg (N4 + polymer) — assumed the same premix"),
}

# Raw-material aliases we infer but Gaurav has not confirmed — applied, and listed as flags.
ASSUMED_RAW_ALIASES = {
    "sn500(kg)": ("sn500", "grease recipes write sn500(kg) — read as SN500 measured by weight (kg basis)"),
}

# Codes that could be one of several things — never resolved automatically.
AMBIGUOUS = {
    "n500vi10kg": "could be sn500vi10kg or w500vi10kg (Gear EP 150–460) — ask Gaurav",
}


def norm(text: str) -> str:
    """Code form: lower case, single spaces, trimmed."""
    return re.sub(r"\s+", " ", str(text)).strip().lower()


def nice(text: str) -> str:
    """Display form: Gaurav's words, spacing tidied, obvious typos fixed."""
    out = re.sub(r"\s+", " ", str(text)).strip()
    for wrong, right in TYPO_FIXES.items():
        out = re.sub(rf"\b{wrong}\b", right, out, flags=re.I)
    return out


def pct_to_bp(value: str) -> int | None:
    """'45.3' → 4530 basis points; '' → None (blank = no amount given)."""
    value = value.strip()
    if value == "":
        return None
    return int(round(float(value) * 100))


def parse_line(cell: str) -> dict:
    """'SAE40/45.3' → {code:'sae40', pctBp:4530}; 'go/' → {code:'go', pctBp:None}; 'brake fluid' → no '/'."""
    text = str(cell).strip()
    if "/" in text:
        name, pct = text.rsplit("/", 1)
        try:
            return {"code": norm(name), "pctBp": pct_to_bp(pct), "raw": text}
        except ValueError:
            pass
    return {"code": norm(text), "pctBp": None, "raw": text, "noPercentSlot": True}


def unit_of(value) -> str | None:
    v = norm(value or "")
    return {"liter": "L", "litre": "L", "l": "L", "kg": "kg", "piece": "piece"}.get(v)


def first_number(value) -> float | None:
    m = re.match(r"\s*(\d+(?:\.\d+)?)", str(value or ""))
    return float(m.group(1)) if m else None


def build(xlsx: Path, out: Path) -> dict:
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    s1, s2, s3 = wb["Sheet1"], wb["Sheet2"], wb["Sheet3"]
    (out / "source").mkdir(parents=True, exist_ok=True)
    for ws, name in ((s1, "sheet1"), (s2, "sheet2"), (s3, "sheet3")):
        with open(out / "source" / f"{name}.csv", "w", newline="") as f:
            w = csv.writer(f)
            for row in ws.iter_rows(values_only=True):
                w.writerow(["" if v is None else v for v in row])

    # ── Sheet1: consumables (packaging) and raw materials ────────────────────
    consumables, raw = [], []
    for r in s1.iter_rows(min_row=2):
        a = r[0].value
        if a in (None, ""):
            continue
        row_no = r[0].row
        name_raw = str(a)
        also = r[1].value
        unit, alt_unit = unit_of(r[2].value), unit_of(r[3].value)
        weight_raw, density = r[4].value, r[5].value
        flags = []
        if norm(name_raw).startswith(("empty ", "beri cap")):
            weight = first_number(weight_raw)
            if weight_raw not in (None, "") and (weight is None or not isinstance(weight_raw, (int, float))):
                flags.append(f'weight written as "{weight_raw}" — read as {weight if weight is not None else "unknown"}')
            code = norm(nice(name_raw))
            consumables.append({
                "code": code,
                "name": nice(name_raw),
                "kind": "consumable",
                "baseUnit": unit or "piece",
                "emptyWeightGrams": weight,
                "aliases": sorted({norm(name_raw)} - {code}),
                "flags": flags,
                "source": f"Sheet1 row {row_no}",
            })
            continue

        # Raw material: "w500/ w600/ hlp/ n100" → code w500, aliases w600, hlp, n100.
        parts = [p.strip() for p in re.split(r"/", name_raw) if p.strip()]
        trade = None
        cleaned = []
        for p in parts:
            m = re.match(r"(.*?)\(\s*trade name\s*\)", p, re.I)
            if m:
                trade = nice(m.group(1)).strip()
                cleaned.append(trade)
            else:
                cleaned.append(p)
        code = norm(cleaned[0])
        aliases = {norm(nice(p)) for p in cleaned[1:]} | {norm(p) for p in cleaned[1:]}
        if norm(code).startswith("grease "):
            aliases.add(code)
            code = code[len("grease "):]
        if also:
            aliases.add(norm(also))
            aliases.add(norm(also).replace(" ", ""))
        if code == "polymer":
            aliases.update({"lubimax141", "lubimax 141"})
        if code == "rpo 710":
            aliases.add("rpo710")
        fixed = norm(nice(code))
        if fixed != code:
            aliases.add(fixed)
        for alias, (target, why) in ASSUMED_RAW_ALIASES.items():
            if target == code:
                aliases.add(alias)
                flags.append(f'alias "{alias}" assumed: {why}')
        aliases.discard(code)
        base = unit or alt_unit
        if unit is None and alt_unit is None:
            flags.append("no unit in the sheet")
        elif unit is None:
            flags.append(f"unit taken from 'another unit' column ({alt_unit})")
        if alt_unit and unit and alt_unit != unit:
            flags.append(f"also measured in {alt_unit}")
        if density in (None, ""):
            flags.append("no density — kg↔L conversion not possible yet")
        if code in ("roll perfect 46",):
            flags.append("also a finished product (ROLL PERFECT 46 (SOLUBLE)) — likely bought in and resold")
        raw.append({
            "code": code,
            "name": nice(name_raw) if not trade else f"{nice(cleaned[0])} ({trade})",
            "kind": "raw_material",
            "baseUnit": base,
            "altUnit": alt_unit if alt_unit != base else None,
            "densityKgPerL": float(density) if isinstance(density, (int, float)) else None,
            "tradeName": trade,
            "aliases": sorted(a for a in aliases if a),
            "flags": flags,
            "source": f"Sheet1 row {row_no}",
        })

    # ── Sheet2: premixes (header says "% in liter" → litre basis) ─────────────
    premixes = []
    for r in s2.iter_rows(min_row=2):
        if r[0].value in (None, ""):
            continue
        code = norm(r[0].value)
        lines = [parse_line(c.value) for c in r[3:] if c.value not in (None, "")]
        # Columns F/H are the kg-basis columns; Gaurav left their amounts blank. Keep only
        # the litre-basis lines (D, E, G) as the recipe; note the kg columns.
        litre_lines = [parse_line(c.value) for c in (r[3], r[4], r[6]) if c.value not in (None, "")]
        premixes.append({
            "code": code,
            "name": str(r[0].value).strip(),
            "tradeName": nice(r[1].value) if r[1].value else None,
            "baseUnit": unit_of(r[2].value) or "L",
            "basis": "litre",
            "lines": litre_lines,
            "kgBasisColumnsBlank": all(l["pctBp"] is None for l in (parse_line(c.value) for c in (r[5], r[7]) if c.value not in (None, ""))),
            "aliases": [],
            "flags": [],
            "source": f"Sheet2 row {r[0].row}",
            "_allCells": [l["raw"] for l in lines],
        })
    for alias, (target, why) in ASSUMED_ALIASES.items():
        for p in premixes:
            if p["code"] == target:
                p["aliases"].append(alias)
                p["flags"].append(f'alias "{alias}" assumed: {why}')

    # ── Sheet3: house base blends ────────────────────────────────────────────
    blends = []
    for r in s3.iter_rows(min_row=2):
        if r[0].value in (None, ""):
            continue
        code = norm(r[0].value)
        aliases = sorted({norm(c.value) for c in r[1:4] if c.value not in (None, "")} - {code})
        lines = [parse_line(c.value) for c in r[4:] if c.value not in (None, "")]
        flags = []
        if not lines:
            flags.append("no recipe — likely a bought-in finished product, not a blend")
        blends.append({
            "code": code,
            "name": nice(r[0].value),
            "aliases": aliases,
            "baseUnit": "L",
            "basis": "unconfirmed (by volume assumed)",
            "lines": lines,
            "flags": flags,
            "source": f"Sheet3 row {r[0].row}",
        })

    # ── Main tab: one recipe per stock pool (product + grade) ────────────────
    main = wb["Product Master (grades)"]
    hdr = [str(c.value).strip() if c.value else "" for c in main[1]]
    col = {h: i for i, h in enumerate(hdr)}

    def txt(v):
        if isinstance(v, float) and v.is_integer():
            return str(int(v))
        return "" if v is None else str(v).strip()

    def slug(s):
        return re.sub(r"[^a-z0-9]+", "-", str(s).lower()).strip("-")

    pools: dict[str, dict] = {}
    for row in main.iter_rows(min_row=2, values_only=True):
        if not any(v not in (None, "") for v in row):
            continue
        prod, grade = txt(row[col["Product"]]), txt(row[col["Grade"]])
        if grade == "—":
            grade = ""
        sku = txt(row[col["SKU"]])
        key = slug(f"{prod} {grade}") if grade else slug(prod)
        recipe = [txt(v) for v in row[5:12] if v not in (None, "")]
        p = pools.setdefault(key, {"productKey": key, "product": prod, "grade": grade, "skus": [], "recipes": {}})
        p["skus"].append(sku)
        p["recipes"].setdefault(tuple(norm(x) for x in recipe), []).append(sku)
        p.setdefault("rawRecipe", recipe)

    product_recipes = []
    for key, p in sorted(pools.items()):
        variants = sorted(p["recipes"].items(), key=lambda kv: -len(kv[1]))
        main_recipe = list(variants[0][0])
        if [norm(x) for x in p["rawRecipe"]] == main_recipe:
            main_recipe = list(p["rawRecipe"])
        flags = []
        if len(variants) > 1:
            flags.append("packs of this pool carry DIFFERENT recipes in the sheet: " +
                         " | ".join(f"{', '.join(skus)} → {' + '.join(r) or '(none)'}" for r, skus in variants))
        if not main_recipe:
            flags.append("no recipe in the sheet")
        product_recipes.append({
            "productKey": key,
            "product": p["product"],
            "grade": p["grade"],
            "skus": p["skus"],
            "basis": "unconfirmed (by volume assumed)",
            "lines": [parse_line(x) for x in main_recipe],
            "flags": flags,
        })

    # ── Resolve every ingredient code ────────────────────────────────────────
    index: dict[str, tuple[str, str]] = {}

    def register(kind, item):
        for c in [item["code"], *item.get("aliases", [])]:
            index.setdefault(c, (kind, item["code"]))
            index.setdefault(c.replace(" ", ""), (kind, item["code"]))

    for x in raw:
        register("raw_material", x)
    for x in premixes:
        register("premix", x)
    for x in blends:
        register("blend", x)

    usage: dict[str, list[str]] = defaultdict(list)

    def resolve_lines(where, lines):
        total, has_blank = 0, False
        for line in lines:
            hit = index.get(line["code"]) or index.get(line["code"].replace(" ", ""))
            line["resolvesTo"] = {"kind": hit[0], "code": hit[1]} if hit else None
            usage[line["code"]].append(where)
            if line["pctBp"] is None:
                has_blank = True
            else:
                total += line["pctBp"]
        return total, has_blank

    for group, items, label in ((premixes, premixes, "premix"), (blends, blends, "blend")):
        for it in items:
            total, blank = resolve_lines(f"{label} {it['code']}", it["lines"])
            it["totalBp"] = total
            if it["lines"] and total != 10000:
                it["flags"].append(f"stated percentages add up to {total / 100:g}%, not 100%")
            if blank:
                it["flags"].append("has ingredients with no % (meaning unconfirmed — read as 'alternative')")
    for pr in product_recipes:
        total, blank = resolve_lines(f"product {pr['productKey']}", pr["lines"])
        pr["totalBp"] = total
        lines = pr["lines"]
        if len(lines) == 1 and lines[0]["pctBp"] in (None, 10000) and lines[0].get("resolvesTo"):
            # One ingredient, all of it: the product is filled straight from that item.
            # For base oils that's normal (White Oil 32 = W150); for ready greases or brake
            # fluid it means bought in and repacked. Gaurav to confirm which (open question).
            target = lines[0]["resolvesTo"]
            pr["singleIngredient"] = target["code"]
            if lines[0]["pctBp"] is None:
                lines[0]["pctBp"] = 10000
                lines[0]["percentAssumed"] = True
            pr["totalBp"] = 10000
            pr["flags"].append(f"one-ingredient recipe: filled straight from {target['code']} ({target['kind']}) — "
                               "for greases / brake fluid this likely means bought in and repacked")
            continue
        if pr["lines"] and total != 10000:
            pr["flags"].append(f"stated percentages add up to {total / 100:g}%, not 100%")
        if blank:
            pr["flags"].append("has ingredients with no % (meaning unconfirmed — read as 'alternative')")
        for line in pr["lines"]:
            if line["resolvesTo"] is None:
                why = AMBIGUOUS.get(line["code"])
                pr["flags"].append(f'ingredient "{line["raw"]}" ' + (f"is ambiguous: {why}" if why else "is not defined in Sheet1/2/3"))

    resolution = {
        "resolved": {c: {"kind": index[c][0], "code": index[c][1], "usedIn": len(w)}
                     for c, w in sorted(usage.items()) if c in index or c.replace(" ", "") in index
                     for _ in [0] if (c in index)},
        "unresolved": {c: {"usedIn": len(w), "examples": sorted(set(w))[:6]}
                       for c, w in sorted(usage.items(), key=lambda kv: -len(kv[1]))
                       if c not in index and c.replace(" ", "") not in index},
        "assumedAliases": {**{a: t for a, (t, _) in ASSUMED_ALIASES.items()},
                           **{a: t for a, (t, _) in ASSUMED_RAW_ALIASES.items()}},
        "ambiguous": AMBIGUOUS,
    }
    for c, w in usage.items():
        if c not in index and c.replace(" ", "") in index:
            resolution["resolved"][c] = {"kind": index[c.replace(" ", "")][0], "code": index[c.replace(" ", "")][1],
                                         "usedIn": len(w), "matchedWithoutSpaces": True}

    for p in premixes:
        p.pop("_allCells", None)

    files = {
        "raw-materials.json": raw,
        "consumables.json": consumables,
        "premixes.json": premixes,
        "blends.json": blends,
        "product-recipes.json": product_recipes,
        "resolution.json": resolution,
    }
    for name, data in files.items():
        (out / name).write_text(json.dumps(data, indent=1, ensure_ascii=False) + "\n")

    return {
        "rawMaterials": len(raw),
        "consumables": len(consumables),
        "premixes": len(premixes),
        "blends": len(blends),
        "productPools": len(product_recipes),
        "poolsWithRecipe": sum(1 for p in product_recipes if p["lines"]),
        "unresolvedCodes": len(resolution["unresolved"]),
    }


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    out_dir = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "office" / "data" / "master"
    print(json.dumps(build(Path(sys.argv[1]), out_dir), indent=1))
